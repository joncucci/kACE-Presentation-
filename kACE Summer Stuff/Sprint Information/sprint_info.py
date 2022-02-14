# Imports
import json, requests, urllib3, sys, os, csv, getpass, errno
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.chart import BarChart, Series, Reference, BarChart3D
from openpyxl.chart.label import DataLabelList
from prettytable import PrettyTable
from datetime import date
from subprocess import Popen
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

JIRA_INSTALLATION = "jira.gfinet.com"
USERNAME, PASSWORD = 'gjapi', 'gjapi'

# Class task for holding information easier
class Task:
    def __init__(self, key, summary, hours, points, rate, resolution):
        self.key = key
        self.summary = summary
        self.hours = hours
        self.points = points
        self.rate = rate
        self.resolution = resolution

    def __repr__(self):
        return repr(self.key)

    def list(self):
        return [self.key, self.summary, self.hours, self.points, self.rate, self.resolution]

# Makes directory path if it doesn't already exist
def make_path(path):

    try:
        os.makedirs(path)
    except OSError as e:
        if e.errno != errno.EEXIST:
            raise

def main(sprint):

    # Grabs system arguments
    sprint = '"' + ' '.join(sprint) + '"'

    # Fetches JSON data
    print("Fetching data...")
    filter = 'sprint = {}'.format(sprint)
    url = "http://%s/rest/api/2/search?jql=%s&fields=key,summary,issuetype,status&maxResults=1000&startAt=%d&os_username=%s&os_password=%s" % (JIRA_INSTALLATION, filter, 0, USERNAME, PASSWORD)
    retrieve = requests.get(url = url,headers={'Content-Type': 'application/json'},auth=(USERNAME, PASSWORD), verify=False)

    # Information Accumulation
    tasks, total_hours, total_points, total_rate, total_rate = [], [], [], [], []

    # Pretty Table setup
    table = PrettyTable()
    table.title = "Tasks for {}".format(sprint)
    fieldnames = ["Task", "Message", "Time (HR)", "Task Points (TP)", "HR / TP", "Resolution"]
    table.field_names = fieldnames
    table._max_width = {"Message": 50}

    try: 
        retrieve.json()['issues']
    except KeyError:
        print("No sprint found.")
        sys.exit()
    
    print("Writing rows...")

    for issue in retrieve.json()['issues']:

        issue_url = issue['self'] + "?expand=changelog&os_username=%s&os_password=%s" % (USERNAME, PASSWORD)
        retrieve_issue = requests.get(url = issue_url,headers={'Content-Type': 'application/json'},auth=(USERNAME, PASSWORD), verify=False)
        retrieved_issue = retrieve_issue.json()

        try :
            hours = round(retrieved_issue['fields']['timespent']/3600,2)
        except TypeError:
            hours = 0

        try:
            if retrieved_issue['fields']['status']['name'] in ['Business Review', 'Closed']:
                resolution = "Resolved"
            else:
                resolution = retrieved_issue['fields']['resolution']['name']      
        except:
            resolution = "Unresolved"
        
        try:
            points = int(retrieved_issue['fields']['customfield_10502'])
            rate = round(hours / points, 2)
        except:
            points, rate = 0, 0

        tasks.append(Task(
                issue['key'],
                retrieved_issue['fields']['summary'],
                hours,
                points,
                rate,
                resolution
        ))

        table.add_row(tasks[-1].list())
        total_hours.append(hours)
        total_points.append(points)
        total_rate.append(rate)

    sum_hours = sum(total_hours)
    sum_points = sum(total_points)
    total_tasks = len(tasks)

    table.add_row(
        [
            "TOTAL",
            "",
            sum_hours,
            sum_points,
            "",
            ""
        ]
    )

    print(table)





    if sum_points == 0:
        print("Points not yet assigned. No further statistics at this time.")
        sys.exit() 


    # Opening and writing CSV / XLSX files
    print("Writing CSV / XLSX...")
    filename = "_".join(sprint[1:-1].split(" "))
    make_path(r"{}\{}".format(os.getcwd(), filename))
    route = r'{}\{}\{}_{}'.format(os.getcwd(), filename, filename, date.today())

    wb = openpyxl.Workbook()

    for sheet in ['NUMBERS', 'TIME STATISTICS', 'TIME QUARTILES','RATE STATISTICS']:
        wb.create_sheet(sheet)
    wb.remove(wb['Sheet'])

    ws = wb['NUMBERS']
    ws.append([])
    ws.append([sprint])
    ws.append(['{} Tasks'.format(len(tasks))])
    ws.append([])
    ws.append(fieldnames)

    for task in tasks:
        ws.append(task.list())

    ws.append([])

    if sum_points == 0:
        print("Points not yet assigned.")
        sys.exit()

    # Adding statistics to CSV file
    ws.append(["TOTAL", "", sum_hours, sum_points, "",""])
    ws.append(["AVERAGE", "", round(sum_hours / total_tasks,2), round(sum_points / total_tasks, 2), round(sum_hours / sum_points,2), ""])
    ws.append(["MEDIAN", "", sorted(total_hours)[int(total_tasks/2)], sorted(total_points)[int(total_tasks/2)], sorted(total_rate)[int(total_tasks/2)], ""])
    ws.append(["MIN", "", min(total_hours), min(total_points), min(total_rate), ""])
    ws.append(["MAX", "", max(total_hours), max(total_points), max(total_rate), ""])
    ws.append(["RANGE", "", max(total_hours) - min(total_hours), max(total_points) - min(total_points), max(total_rate) - min(total_rate)])
    ws.append(["0% QUARTILE", "", np.percentile(total_hours, 0), np.percentile(total_points, 0), np.percentile(total_rate, 0)])        
    ws.append(["25% QUARTILE", "", np.percentile(total_hours, 25), np.percentile(total_points, 25), np.percentile(total_rate, 25)])
    ws.append(["50% QUARTILE", "", np.percentile(total_hours, 50), np.percentile(total_points, 50), np.percentile(total_rate, 50)])
    ws.append(["75% QUARTILE", "", np.percentile(total_hours, 75), np.percentile(total_points, 75), np.percentile(total_rate, 75)])
    ws.append(["100% QUARTILE", "", np.percentile(total_hours, 100), np.percentile(total_points, 100), np.percentile(total_rate, 100)]) 
    ws.append([])

    # Separates data based on story points
    for total in [True, False]:

        average, median, min_, max_, range_, q0, q1, q2, q3, q4 = ['AVERAGE'], ['MEDIAN'], ['MIN'], ['MAX'], ['RANGE'], ['0% QUARTILE'], ['25% QUARTILE'], ['50% QUARTILE'], ['75% QUARTILE'], ['100% QUARTILE']

        if total:
            header = ["HOURS \ TASK"]
        
        else:
            header = ["RATE \ TASK"]


        for num in range(0,22):

            if not total and num == 0:
                continue
            
            total_hours, p_tasks, total_points, total_rate = [], [], [], []

            for task in tasks:
                if task.points == num:
                    p_tasks.append(task)
                    total_hours.append(task.hours)

            sum_hours, sum_points, total_tasks = sum(total_hours), sum(total_points), len(p_tasks)
                    
            if total_tasks > 0:

                header.append("{}SP ({})".format(num, total_tasks))
                if total:
                    num = 1
                average.append(round((sum_hours / total_tasks)/num, 2))
                median.append(round(sorted(total_hours)[int(total_tasks/2)]/num,2))
                min_.append(round(min(total_hours)/num,2))
                max_.append(round(max(total_hours)/num,2))
                range_.append(round(max(total_hours)/num - min(total_hours)/num,2))
                q0.append(round(np.percentile(total_hours, 0)/num,2))
                q1.append(round(np.percentile(total_hours, 25)/num,2))
                q2.append(round(np.percentile(total_hours, 50)/num,2))
                q3.append(round(np.percentile(total_hours, 75)/num,2))
                q4.append(round(np.percentile(total_hours, 100)/num,2))

        for row in [header, average, median, min_, max_, range_, q0, q1, q2, q3, q4]:
            ws.append(row)
        ws.append([])

    wb.save(route + '.xlsx')
    pd.read_excel(route + '.xlsx').to_csv(route + '.csv', index = None, header = False)
    print("CSV written.")

    '''
    
    Start of graphing
    CHART 1: TIME STATISTICS

    '''
    chart1 = BarChart()
    chart1.type = "col"
    chart1.style = 10
    chart1.title = "TIME STATISTICS"
    chart1.y_axis.title = "HOURS"
    chart1.x_axis.title = "STATISTICS"
    data = Reference(ws, min_col = 2, min_row = 19 + len(tasks), max_col = 1 + len(header), max_row = 24 + len(tasks))
    cats = Reference(ws, min_col = 1, min_row = 20 + len(tasks), max_row = 24 + len(tasks))
    chart1.add_data(data, titles_from_data = True)
    chart1.set_categories(cats)
    chart1.height = 15
    chart1.width = 25
    chart1.shape = 4
    chart1.dataLabels = DataLabelList()
    chart1.dataLabels.showVal = True

    wb['TIME STATISTICS'].add_chart(chart1, 'B2') # "G" + str(14 + len(tasks)))


    '''

    CHART 2: TIME QUARTILES
    
    '''

    chart2 = BarChart()
    chart2.type = "col"
    chart2.style = 10
    chart2.title = "TIME QUARTILES"
    chart2.y_axis.title = "HOURS"
    chart2.x_axis.title = "QUARTILES"
    sp = Reference(ws, min_col = 2, min_row = 25 + len(tasks), max_col = 1 + len(header))
    data = Reference(ws, min_col = 2, min_row = 26 + len(tasks), max_col = 1 + len(header), max_row = 29 + len(tasks))
    cats = Reference(ws, min_col = 1, min_row = 13 + len(tasks), max_row = 19 + len(tasks))
    chart2.add_data(data, titles_from_data = True)
    chart2.set_categories(cats)
    chart2.height = 15
    chart2.width = 25
    chart2.shape = 4
    chart2.dataLabels = DataLabelList()
    chart2.dataLabels.showVal = True

    wb['TIME QUARTILES'].add_chart(chart2, 'B2')   
    
    wb.save(route + '.xlsx')
    print("XLSX written.")

    file_path = r"{}\{}".format(os.getcwd(), filename)
    print("Copies with more statistics has been saved to:\n{}\n".format(file_path))

    # Pop opens the XLSX file in excel
    os.startfile(route + '.xlsx')




if __name__ == "__main__":

    if len(sys.argv) == 1:
        print("\nEnter a compatible Sprint Name:\nEX.) sprint_info 'Firefly MDP Sprint 7'\n")
    else:
        main(sys.argv[1:]) 